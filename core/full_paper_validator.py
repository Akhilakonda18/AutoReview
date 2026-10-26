"""Comprehensive pre-submission screening for research papers.

Checks are heuristic screening aids, not a substitute for author review or
an official venue's submission validator. Supports PDF and DOCX inputs.
"""

from __future__ import annotations

import json
import re
from collections import Counter
from datetime import datetime
from pathlib import Path

try:
    import fitz
except ImportError:
    fitz = None


SECTION_ALIASES = {
    "abstract": ("abstract", "summary"),
    "introduction": ("introduction",),
    "related work": (
        "related work",
        "related works",
        "literature review",
        "literature survey",
        "background",
    ),
    "methodology": (
        "methodology",
        "methods",
        "materials and methods",
        "proposed method",
        "proposed approach",
    ),
    "results": (
        "results",
        "experimental results",
        "results and discussion",
        "evaluation",
    ),
    "discussion": ("discussion",),
    "limitations": ("limitations", "threats to validity"),
    "conclusion": (
        "conclusion",
        "conclusions",
        "conclusion and future work",
        "conclusions and future work",
    ),
    "references": ("references", "bibliography", "works cited"),
}

PLACEHOLDER_PATTERNS = [
    r"\bTODO\b",
    r"\bTBD\b",
    r"\bPLACEHOLDER\b",
    r"\[insert\b",
    r"\[your\s+\w+",
    r"\bXXX+\b",
    r"\bLorem ipsum\b",
    r"Figure\s+\?",
    r"Table\s+\?",
    r"\[citation needed\]",
]


def _result(
    check: str,
    status: str,
    details: str,
    location: str = "",
) -> dict:
    """Create a standardized validation finding."""
    return {
        "Check": check,
        "Status": status,
        "Details": details,
        "Location": location,
    }


def _read_pdf(path: Path) -> dict:
    """Extract text and page dimensions from a PDF."""
    if fitz is None:
        raise RuntimeError(
            "PDF checks require PyMuPDF. Install it using: pip install pymupdf"
        )

    pages = []
    page_info = []

    with fitz.open(str(path)) as doc:
        if doc.is_encrypted:
            raise ValueError(
                "The PDF is password-protected or encrypted."
            )

        for index, page in enumerate(doc):
            text = page.get_text("text") or ""
            pages.append(text)

            page_info.append({
                "page": index + 1,
                "width_pt": round(float(page.rect.width), 2),
                "height_pt": round(float(page.rect.height), 2),
                "text_chars": len(text.strip()),
                "is_blank_or_unextractable": not bool(text.strip()),
            })

    return {
        "text": "\n".join(pages),
        "pages": page_info,
        "format": "PDF",
    }


def _read_docx(path: Path) -> dict:
    """Extract document text and paragraph-formatting properties."""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "DOCX checks require python-docx. "
            "Install it using: pip install python-docx"
        ) from exc

    doc = Document(str(path))
    parts = []
    paragraphs = []

    for index, para in enumerate(doc.paragraphs, start=1):
        txt = para.text or ""
        parts.append(txt)

        if txt.strip():
            fmt = para.paragraph_format

            paragraphs.append({
                "index": index,
                "text": txt,
                "style": (
                    getattr(para.style, "name", "")
                    if para.style
                    else ""
                ),
                "left_indent_pt": (
                    round(fmt.left_indent.pt, 2)
                    if fmt.left_indent is not None
                    else None
                ),
                "first_line_indent_pt": (
                    round(fmt.first_line_indent.pt, 2)
                    if fmt.first_line_indent is not None
                    else None
                ),
                "space_before_pt": (
                    round(fmt.space_before.pt, 2)
                    if fmt.space_before is not None
                    else None
                ),
                "space_after_pt": (
                    round(fmt.space_after.pt, 2)
                    if fmt.space_after is not None
                    else None
                ),
                "line_spacing": fmt.line_spacing,
            })

    for table in doc.tables:
        for row in table.rows:
            parts.append(
                " | ".join(cell.text for cell in row.cells)
            )

    return {
        "text": "\n".join(parts),
        "pages": [],
        "format": "DOCX",
        "paragraphs": paragraphs,
        "table_count": len(doc.tables),
        "inline_shape_count": len(doc.inline_shapes),
    }


def _find_sections(text: str) -> set[str]:
    """Identify common research-paper section headings."""
    detected = set()

    for line in text.splitlines():
        heading = re.sub(
            r"\s+",
            " ",
            line.strip(),
        ).strip(" .:-").lower()

        heading = re.sub(
            r"^(?:(?:\d+(?:\.\d+)*|[ivx]+)[.)]?\s+)",
            "",
            heading,
        )

        for canonical, aliases in SECTION_ALIASES.items():
            if heading in aliases:
                detected.add(canonical)

    # Fallback for headings that PDF extraction may format differently.
    for canonical, aliases in SECTION_ALIASES.items():
        for alias in aliases:
            pattern = (
                r"(?im)^\s*"
                r"(?:(?:\d+(?:\.\d+)*|[IVX]+)[.)]?\s+)?"
                + re.escape(alias)
                + r"\s*[:.]?\s*$"
            )

            if re.search(pattern, text):
                detected.add(canonical)
                break

    return detected


def _line_number(text: str, match_start: int) -> int:
    """Return the approximate line number of a text match."""
    return text.count("\n", 0, match_start) + 1


def validate_research_paper(
    file_path: str | Path,
    required_sections: list[str] | None = None,
    min_pages: int | None = None,
    max_pages: int | None = None,
    expected_page_size: str | None = None,
    check_grammar: bool = True,
) -> dict:
    """
    Screen a research paper for common pre-submission issues.

    Supported formats:
        PDF
        DOCX

    Checks include:
        File readability
        Blank or unextractable PDF pages
        Page count
        Page dimensions
        Required section headings
        Unresolved placeholders
        Spacing and repeated words
        DOCX indentation consistency
        DOCX line-spacing consistency
        References section
        Numeric citation/reference matching
        Figure captions and references
        Table captions and references
        Optional grammar and spelling screening
    """

    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(
            f"Paper file not found: {path}"
        )

    suffix = path.suffix.lower()

    if suffix == ".pdf":
        extracted = _read_pdf(path)

    elif suffix == ".docx":
        extracted = _read_docx(path)

    else:
        raise ValueError(
            "Supported paper formats are PDF and DOCX."
        )

    text = extracted["text"]
    normalized = re.sub(r"\s+", " ", text).strip()

    checks = []

    # ---------------------------------------------------------
    # 1. File readability
    # ---------------------------------------------------------

    checks.append(
        _result(
            "File readable",
            "PASS" if normalized else "FAIL",
            (
                f"Text extracted from {path.name}."
                if normalized
                else (
                    "No text could be extracted. "
                    "Scanned PDFs may require OCR."
                )
            ),
        )
    )

    # ---------------------------------------------------------
    # 2. PDF page integrity and page count
    # ---------------------------------------------------------

    if suffix == ".pdf":
        page_count = len(extracted["pages"])

        blank_pages = [
            page["page"]
            for page in extracted["pages"]
            if page["is_blank_or_unextractable"]
        ]

        checks.append(
            _result(
                "PDF page integrity",
                "WARNING" if blank_pages else "PASS",
                (
                    f"{page_count} pages detected; "
                    "pages with no extractable text: "
                    f"{blank_pages or 'none'}."
                ),
                (
                    f"Pages {blank_pages}"
                    if blank_pages
                    else ""
                ),
            )
        )

        if min_pages is not None or max_pages is not None:
            lower_ok = (
                min_pages is None
                or page_count >= min_pages
            )

            upper_ok = (
                max_pages is None
                or page_count <= max_pages
            )

            status = (
                "PASS"
                if lower_ok and upper_ok
                else "FAIL"
            )

            checks.append(
                _result(
                    "Page count limit",
                    status,
                    (
                        f"{page_count} pages; configured range "
                        f"{min_pages if min_pages is not None else 'no minimum'}"
                        "–"
                        f"{max_pages if max_pages is not None else 'no maximum'}."
                    ),
                )
            )

        # -----------------------------------------------------
        # 3. Page size
        # -----------------------------------------------------

        if expected_page_size:
            target = (
                expected_page_size.lower().replace(" ", "")
            )

            first = (
                extracted["pages"][0]
                if extracted["pages"]
                else {}
            )

            width = first.get("width_pt", 0)
            height = first.get("height_pt", 0)

            if target == "a4":
                ok = (
                    abs(min(width, height) - 595.28) < 8
                    and abs(max(width, height) - 841.89) < 8
                )

            elif target in {"usletter", "letter"}:
                ok = (
                    abs(min(width, height) - 612) < 8
                    and abs(max(width, height) - 792) < 8
                )

            else:
                ok = None

            checks.append(
                _result(
                    "Page size",
                    (
                        "PASS"
                        if ok is True
                        else "FAIL"
                        if ok is False
                        else "UNVERIFIED"
                    ),
                    (
                        f"First page is {width:.0f} × "
                        f"{height:.0f} pt; configured target: "
                        f"{expected_page_size}."
                    ),
                )
            )

    else:
        checks.append(
            _result(
                "Page count limit",
                "UNVERIFIED",
                (
                    "DOCX page count depends on Word or "
                    "LibreOffice rendering. Export to PDF "
                    "for a reliable page count."
                ),
            )
        )

    # ---------------------------------------------------------
    # 4. Required sections
    # ---------------------------------------------------------

    detected = _find_sections(text)

    required = required_sections or [
        "Abstract",
        "Introduction",
        "Methodology",
        "Results",
        "Conclusion",
        "References",
    ]

    missing = []

    for name in required:
        key = name.strip().lower()

        aliases = SECTION_ALIASES.get(
            key,
            (key,),
        )

        found = False

        for alias in aliases:
            pattern = (
                r"(?im)^\s*"
                r"(?:(?:\d+(?:\.\d+)*|[IVX]+)[.)]?\s+)?"
                + re.escape(alias)
                + r"\s*[:.]?\s*$"
            )

            if re.search(pattern, text):
                found = True
                break

        if not found:
            missing.append(name)

    checks.append(
        _result(
            "Required sections",
            "FAIL" if missing else "PASS",
            (
                "Missing or not detected: "
                + ", ".join(missing)
                if missing
                else (
                    "All configured required section "
                    "headings were detected."
                )
            ),
        )
    )

    # ---------------------------------------------------------
    # 5. Unresolved placeholders
    # ---------------------------------------------------------

    placeholders = []

    for pattern in PLACEHOLDER_PATTERNS:
        for match in re.finditer(
            pattern,
            text,
            re.IGNORECASE,
        ):
            placeholders.append({
                "text": match.group(0),
                "line": _line_number(
                    text,
                    match.start(),
                ),
            })

    checks.append(
        _result(
            "Unresolved placeholders",
            "WARNING" if placeholders else "PASS",
            (
                f"Potential placeholders found: "
                f"{len(placeholders)}."
                if placeholders
                else (
                    "No common placeholder patterns "
                    "were detected."
                )
            ),
            "; ".join(
                f"{item['text']} (line {item['line']})"
                for item in placeholders[:10]
            ),
        )
    )

    # ---------------------------------------------------------
    # 6. Spacing and typographical patterns
    # ---------------------------------------------------------

    style_issues = []

    patterns = [
        (
            r"[ \t]+[,.!?;:]",
            "space before punctuation",
        ),
        (
            r"\b([A-Za-z]{2,})\s+\1\b",
            "repeated word",
        ),
        (
            r"\s{3,}",
            "three or more consecutive spaces",
        ),
        (
            r"(?m)^\s+$",
            "whitespace-only line",
        ),
    ]

    for pattern, description in patterns:
        for match in re.finditer(
            pattern,
            text,
            re.IGNORECASE,
        ):
            style_issues.append({
                "issue": description,
                "text": (
                    match.group(0)[:40]
                    .replace("\n", "\\n")
                ),
                "line": _line_number(
                    text,
                    match.start(),
                ),
            })

    checks.append(
        _result(
            "Spacing and typographical patterns",
            "WARNING" if style_issues else "PASS",
            (
                f"{len(style_issues)} potential "
                "spacing/typographical issue(s) detected."
                if style_issues
                else (
                    "No common spacing patterns "
                    "were detected."
                )
            ),
            "; ".join(
                f"{item['issue']} (line {item['line']})"
                for item in style_issues[:12]
            ),
        )
    )

    # ---------------------------------------------------------
    # 7. DOCX indentation and line spacing
    # ---------------------------------------------------------

    if suffix == ".docx":
        paragraphs = extracted.get(
            "paragraphs",
            [],
        )

        body_paragraphs = [
            paragraph
            for paragraph in paragraphs
            if paragraph["text"].strip()
            and not re.match(
                r"^(abstract|references|introduction|conclusion)\s*$",
                paragraph["text"].strip(),
                re.IGNORECASE,
            )
        ]

        first_indents = [
            paragraph["first_line_indent_pt"]
            for paragraph in body_paragraphs
            if paragraph["first_line_indent_pt"] is not None
        ]

        if first_indents:
            rounded = Counter(
                round(value, 1)
                for value in first_indents
            )

            common_indent, _ = rounded.most_common(1)[0]

            inconsistent = [
                paragraph
                for paragraph in body_paragraphs
                if (
                    paragraph["first_line_indent_pt"] is not None
                    and abs(
                        paragraph["first_line_indent_pt"]
                        - common_indent
                    ) > 3
                )
            ]

            checks.append(
                _result(
                    "Paragraph indentation consistency",
                    "WARNING" if inconsistent else "PASS",
                    (
                        f"Most common first-line indent: "
                        f"{common_indent:.1f} pt; "
                        f"{len(inconsistent)} paragraph(s) "
                        "differ by more than 3 pt."
                    ),
                    ", ".join(
                        f"paragraph {paragraph['index']}"
                        for paragraph in inconsistent[:15]
                    ),
                )
            )

        else:
            checks.append(
                _result(
                    "Paragraph indentation consistency",
                    "UNVERIFIED",
                    (
                        "No explicit first-line indentation "
                        "was found in the DOCX paragraph "
                        "properties. Compare with the venue template."
                    ),
                )
            )

        line_spacings = [
            paragraph["line_spacing"]
            for paragraph in body_paragraphs
            if isinstance(
                paragraph["line_spacing"],
                (int, float),
            )
        ]

        if line_spacings:
            rounded = Counter(
                round(float(value), 2)
                for value in line_spacings
            )

            common_spacing, _ = rounded.most_common(1)[0]

            odd = sum(
                1
                for value in line_spacings
                if abs(
                    float(value) - common_spacing
                ) > 0.1
            )

            checks.append(
                _result(
                    "Paragraph line-spacing consistency",
                    "WARNING" if odd else "PASS",
                    (
                        f"Most common line-spacing setting: "
                        f"{common_spacing}; {odd} paragraph(s) "
                        "differ. Exact venue rules still need confirmation."
                    ),
                )
            )

        else:
            checks.append(
                _result(
                    "Paragraph line-spacing consistency",
                    "UNVERIFIED",
                    (
                        "Line spacing may be inherited from "
                        "styles and could not be reliably inferred "
                        "from direct paragraph properties."
                    ),
                )
            )

    else:
        checks.append(
            _result(
                "Indentation and editable styles",
                "UNVERIFIED",
                (
                    "PDF text extraction cannot reliably "
                    "distinguish intentional indentation from "
                    "layout effects. Upload DOCX for paragraph-style screening."
                ),
            )
        )

    # ---------------------------------------------------------
    # 8. References section
    # ---------------------------------------------------------

    refs_match = re.search(
        r"(?im)^\s*(?:\d+(?:\.\d+)*[.)]?\s*)?"
        r"(references|bibliography|works cited)\s*$",
        text,
    )

    references_text = (
        text[refs_match.end():]
        if refs_match
        else ""
    )

    body_text = (
        text[:refs_match.start()]
        if refs_match
        else text
    )

    references_present = bool(
        refs_match
        and references_text.strip()
    )

    checks.append(
        _result(
            "References section",
            "PASS" if references_present else "FAIL",
            (
                "A references heading and following "
                "text were detected."
                if references_present
                else (
                    "A references/bibliography heading "
                    "with content was not detected."
                )
            ),
        )
    )

    # ---------------------------------------------------------
    # 9. Numeric citation/reference matching
    # ---------------------------------------------------------

    numeric_citations = set()

    for match in re.finditer(
        r"\[(\d{1,3})(?:\s*[-,]\s*\d{1,3})*\]",
        body_text,
    ):
        numeric_citations.add(
            int(match.group(1))
        )

    numeric_reference_labels = set()

    for match in re.finditer(
        r"(?m)^\s*\[(\d{1,3})\]\s+",
        references_text,
    ):
        numeric_reference_labels.add(
            int(match.group(1))
        )

    if numeric_citations or numeric_reference_labels:
        uncited = sorted(
            numeric_reference_labels - numeric_citations
        )

        missing_references = sorted(
            numeric_citations - numeric_reference_labels
        )

        status = (
            "WARNING"
            if uncited or missing_references
            else "PASS"
        )

        checks.append(
            _result(
                "Numeric citation/reference cross-check",
                status,
                (
                    f"{len(numeric_citations)} cited number(s); "
                    f"{len(numeric_reference_labels)} numbered "
                    "reference(s). "
                    f"Potential uncited references: "
                    f"{uncited[:15] or 'none'}; "
                    "citations without matching numbered "
                    f"reference: {missing_references[:15] or 'none'}."
                ),
                (
                    "Approximate pattern matching; ranges "
                    "and complex citation styles may be missed."
                ),
            )
        )

    else:
        checks.append(
            _result(
                "Citation/reference cross-check",
                "UNVERIFIED",
                (
                    "No clear numeric citation pattern "
                    "was detected. Author-year and venue-specific "
                    "styles require additional checking."
                ),
            )
        )

    # ---------------------------------------------------------
    # 10. Figure captions and references
    # ---------------------------------------------------------

    figure_mentions = len(
        re.findall(
            r"\bfig(?:ure)?\.?\s*\d+\b",
            body_text,
            re.IGNORECASE,
        )
    )

    figure_captions = re.findall(
        r"(?im)^\s*(?:fig(?:ure)?\.?)\s*\d+\s*[:. -]",
        text,
    )

    figure_status = (
        "WARNING"
        if figure_mentions and not figure_captions
        else "PASS"
        if figure_captions or not figure_mentions
        else "UNVERIFIED"
    )

    checks.append(
        _result(
            "Figure captions and references",
            figure_status,
            (
                f"Body figure mentions: {figure_mentions}; "
                f"caption-like lines: {len(figure_captions)}. "
                "Verify every figure is numbered, captioned, "
                "readable, and referenced."
            ),
        )
    )

    # ---------------------------------------------------------
    # 11. Table captions and references
    # ---------------------------------------------------------

    table_mentions = len(
        re.findall(
            r"\btables?\s*\d+\b",
            body_text,
            re.IGNORECASE,
        )
    )

    table_captions = re.findall(
        r"(?im)^\s*table\s+\d+\s*[:. -]",
        text,
    )

    table_status = (
        "WARNING"
        if table_mentions and not table_captions
        else "PASS"
        if table_captions or not table_mentions
        else "UNVERIFIED"
    )

    checks.append(
        _result(
            "Table captions and references",
            table_status,
            (
                f"Body table mentions: {table_mentions}; "
                f"caption-like lines: {len(table_captions)}. "
                "Verify numbering, captions, and in-text references."
            ),
        )
    )

    # ---------------------------------------------------------
    # 12. Optional grammar and spelling screening
    # ---------------------------------------------------------

    if check_grammar:
        try:
            import language_tool_python

            tool = language_tool_python.LanguageTool(
                "en-US"
            )

            matches = tool.check(
                text[:100000]
            )

            tool.close()

            checks.append(
                _result(
                    "Grammar and spelling screening",
                    "WARNING" if matches else "PASS",
                    (
                        f"LanguageTool reported "
                        f"{len(matches)} potential issue(s)."
                        if matches
                        else (
                            "No issues reported by LanguageTool. "
                            "Automated grammar checks can miss errors."
                        )
                    ),
                    "; ".join(
                        (
                            f"offset {getattr(match, 'offset', '?')}: "
                            f"{match.message}"
                        )
                        for match in matches[:12]
                    ),
                )
            )

        except Exception as exc:
            checks.append(
                _result(
                    "Grammar and spelling screening",
                    "UNVERIFIED",
                    (
                        "Optional LanguageTool check unavailable. "
                        "Install language-tool-python and ensure "
                        "its Java/runtime requirements are met."
                    ),
                    str(exc)[:250],
                )
            )

    else:
        checks.append(
            _result(
                "Grammar and spelling screening",
                "UNVERIFIED",
                "Grammar screening was disabled.",
            )
        )

    # ---------------------------------------------------------
    # 13. Calculate overall status
    # ---------------------------------------------------------

    statuses = [
        check["Status"]
        for check in checks
    ]

    counts = {
        key: statuses.count(key)
        for key in (
            "PASS",
            "FAIL",
            "WARNING",
            "UNVERIFIED",
        )
    }

    if counts["FAIL"]:
        overall = "ISSUES FOUND"

    elif counts["WARNING"] or counts["UNVERIFIED"]:
        overall = "REVIEW REQUIRED"

    else:
        overall = "NO ISSUES DETECTED IN CHECKED ITEMS"

    return {
        "paper": path.name,
        "file_format": extracted["format"],
        "checked_at": datetime.now().isoformat(
            timespec="seconds"
        ),
        "overall_status": overall,
        "summary": counts,
        "detected_sections": sorted(detected),
        "missing_sections": missing,
        "findings": checks,
        "notes": [
            (
                "Automated screening does not prove scientific "
                "correctness, novelty, ethical approval, or acceptance."
            ),
            (
                "Formatting checks depend on the official guidelines "
                "and template supplied by the author."
            ),
            (
                "PDF indentation, visual alignment, equations, "
                "and figure readability should also be checked visually."
            ),
        ],
    }


def save_report(
    result: dict,
    output_path: str | Path,
) -> Path:
    """Save validation findings as a JSON report."""
    path = Path(output_path)

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    path.write_text(
        json.dumps(
            result,
            indent=2,
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    return path